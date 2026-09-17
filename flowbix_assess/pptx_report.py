"""Generates the final PPTX presentation from the same Finding list used by
the HTML report — this is the deliverable described in the "Automação
Assessment" flow: after both Zabbix (via API) and Infraestrutura (via the
SSH probe) have been validated and merged, the automation reads that data
and builds a ready-to-present deck.

This is a generic, brand-neutral template (dark navy + accent colors,
matching the HTML report's palette). If Flowbix has an official .pptx
template with logo/branding, swap PLACEHOLDER content here to load it via
`Presentation(existing_template_path)` instead of `Presentation()` — the
slide-building functions below only need slide_layouts[6] (blank) to exist.
"""
from __future__ import annotations

from collections import Counter, defaultdict

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .models import Severity

NAVY = RGBColor(0x0F, 0x17, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_MUTED = RGBColor(0xB6, 0xC2, 0xD9)
MUTED = RGBColor(0x66, 0x70, 0x85)
CRITICAL = RGBColor(0xB4, 0x23, 0x18)
WARNING = RGBColor(0xB5, 0x47, 0x08)
INFO = RGBColor(0x17, 0x5C, 0xD3)
MANUAL = RGBColor(0x69, 0x41, 0xC6)

SEVERITY_COLOR = {
    Severity.CRITICAL: CRITICAL,
    Severity.WARNING: WARNING,
    Severity.INFO: INFO,
    Severity.MANUAL_REVIEW: MANUAL,
}
SEVERITY_LABEL = {
    Severity.CRITICAL: "CRÍTICO",
    Severity.WARNING: "ATENÇÃO",
    Severity.INFO: "INFORMATIVO",
    Severity.MANUAL_REVIEW: "REVISÃO MANUAL",
}

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
BLANK_LAYOUT = 6


def _textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height).text_frame


def _background(slide, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _set_text(paragraph, text, size, color, bold=False, align=None):
    paragraph.text = text
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = color
    paragraph.font.bold = bold
    if align:
        paragraph.alignment = align


def _title_slide(prs, config):
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    _background(slide, NAVY)
    tf = _textbox(slide, Inches(0.9), Inches(2.8), Inches(11.5), Inches(1.5))
    _set_text(tf.paragraphs[0], "Assessment de Infraestrutura e Zabbix", 40, WHITE, bold=True)
    tf2 = _textbox(slide, Inches(0.9), Inches(4.1), Inches(11.5), Inches(1))
    hosting = (config.client.get("hosting_type") or "").strip()
    hosting_label = {"cloud": " · Ambiente em nuvem", "on-premise": " · Ambiente on-premise",
                      "hybrid": " · Ambiente híbrido"}.get(hosting, "")
    _set_text(tf2.paragraphs[0],
              f"{config.client_name} · Versão-alvo Zabbix {config.target_zabbix_version}{hosting_label}",
              18, LIGHT_MUTED)
    return slide


def _summary_slide(prs, findings):
    counts = Counter(f.severity for f in findings)
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    title_tf = _textbox(slide, Inches(0.7), Inches(0.5), Inches(12), Inches(0.9))
    _set_text(title_tf.paragraphs[0], "Resumo Executivo", 30, NAVY, bold=True)

    cards = [
        ("Total de achados", len(findings), NAVY),
        ("Críticos", counts.get(Severity.CRITICAL, 0), CRITICAL),
        ("Atenção", counts.get(Severity.WARNING, 0), WARNING),
        ("Informativos", counts.get(Severity.INFO, 0), INFO),
        ("Revisão Manual", counts.get(Severity.MANUAL_REVIEW, 0), MANUAL),
    ]
    card_w = Inches(2.2)
    gap = Inches(0.3)
    left = Inches(0.7)
    for label, value, color in cards:
        tf = _textbox(slide, left, Inches(2.2), card_w, Inches(1.8))
        p1 = tf.paragraphs[0]
        _set_text(p1, str(value), 44, color, bold=True, align=PP_ALIGN.CENTER)
        p2 = tf.add_paragraph()
        _set_text(p2, label, 13, MUTED, align=PP_ALIGN.CENTER)
        left += card_w + gap
    return slide


def _agenda_slide(prs, sections):
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    title_tf = _textbox(slide, Inches(0.7), Inches(0.5), Inches(12), Inches(0.9))
    _set_text(title_tf.paragraphs[0], "Tópicos Avaliados", 30, NAVY, bold=True)

    tf = _textbox(slide, Inches(0.9), Inches(1.8), Inches(11), Inches(5))
    for i, section in enumerate(sections):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set_text(p, f"•  {section}", 22, NAVY)
        p.space_after = Pt(14)
    return slide


def _finding_slide(prs, section, finding):
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])

    badge_tf = _textbox(slide, Inches(0.7), Inches(0.4), Inches(4), Inches(0.5))
    _set_text(badge_tf.paragraphs[0], SEVERITY_LABEL[finding.severity], 13,
              SEVERITY_COLOR[finding.severity], bold=True)

    section_tf = _textbox(slide, Inches(8.3), Inches(0.4), Inches(4.3), Inches(0.5))
    _set_text(section_tf.paragraphs[0], section, 13, MUTED, align=PP_ALIGN.RIGHT)

    title_tf = _textbox(slide, Inches(0.7), Inches(1.0), Inches(11.9), Inches(1.1))
    _set_text(title_tf.paragraphs[0], finding.title, 26, NAVY, bold=True)

    body_tf = _textbox(slide, Inches(0.7), Inches(2.4), Inches(11.9), Inches(2.4))
    _set_text(body_tf.paragraphs[0], finding.description, 18, NAVY)

    rec_tf = _textbox(slide, Inches(0.7), Inches(5.1), Inches(11.9), Inches(2))
    _set_text(rec_tf.paragraphs[0], "Recomendação", 14, MUTED, bold=True)
    rec_body = rec_tf.add_paragraph()
    _set_text(rec_body, finding.recommendation, 18, NAVY)
    return slide


def _closing_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
    _background(slide, NAVY)
    tf = _textbox(slide, Inches(0.9), Inches(3.2), Inches(11.5), Inches(1.5))
    _set_text(tf.paragraphs[0], "Fale conosco", 34, WHITE, bold=True)
    p2 = tf.add_paragraph()
    _set_text(p2, "flowbix.com", 16, LIGHT_MUTED)
    return slide


def render(config, findings, output_path: str) -> str:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    by_section = defaultdict(list)
    for finding in findings:
        by_section[finding.section].append(finding)

    _title_slide(prs, config)
    _summary_slide(prs, findings)
    _agenda_slide(prs, list(by_section.keys()))

    for section, section_findings in by_section.items():
        for finding in section_findings:
            _finding_slide(prs, section, finding)

    _closing_slide(prs)

    prs.save(output_path)
    return output_path
