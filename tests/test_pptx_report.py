"""Regression tests for flowbix_assess.pptx_report.

Guards the text-overflow bug found when real-world long titles/
descriptions were used: table cells have a fixed pixel height in PPTX (no
reflow like HTML), so unbounded text silently bled into the next row.
"""
from __future__ import annotations

from datetime import datetime

from pptx import Presentation

from flowbix_assess import pptx_report as pr
from flowbix_assess.models import Finding, Severity

LONG_TITLE = "Itens de baixa prioridade com intervalo curto e sem trigger em 999 template(s) muito compridos de verdade"
LONG_DESCRIPTION = (
    "142 item(ns) com coleta abaixo de 60s e sem nenhuma trigger associada — indica menor "
    "criticidade. Templates afetados: 3001_Template_PDV_ER (12), 3300_Grupo_PDV_LABORATORIOS (8), "
    "7300_Template_Cisco_Meraki_Device_MX_ER_GBTech (4), e mais 20 template(s) que ainda nem "
    "foram listados aqui para deixar esse texto ainda mais comprido."
)
LONG_RECOMMENDATION = (
    "Usar abordagem de item master com itens dependentes para reduzir a quantidade de conexões "
    "diretas ao banco, considerando o volume de proxies e scripts conectando simultaneamente ao "
    "servidor de banco de dados principal em todos os ambientes avaliados neste assessment."
)


def _text_of(shape) -> str:
    return shape.text_frame.text if shape.has_text_frame else ""


def test_truncate_adds_ellipsis_only_when_needed():
    assert pr._truncate("short text", 50) == "short text"
    truncated = pr._truncate("a" * 300, 50)
    assert len(truncated) == 50
    assert truncated.endswith("…")


def test_long_title_and_description_never_exceed_the_calibrated_budget():
    finding = Finding(
        section="Análise de Templates",
        title=LONG_TITLE,
        severity=Severity.WARNING,
        description=LONG_DESCRIPTION,
        recommendation=LONG_RECOMMENDATION,
        deadline_days=30,
    )

    prs = Presentation()
    prs.slide_width = pr.SLIDE_W
    prs.slide_height = pr.SLIDE_H
    pr._section_slides(prs, "Análise de Templates", [finding], datetime.now())

    slide = list(prs.slides)[0]
    texts = [_text_of(s) for s in slide.shapes]

    title_text = next(t for t in texts if t.startswith("Itens de baixa prioridade"))
    desc_text = next(t for t in texts if t.startswith("142 item(ns)"))
    rec_text = next(t for t in texts if "Usar abordagem" in t)

    assert len(title_text) <= pr.TITLE_MAX_CHARS
    assert len(desc_text) <= pr.DESCRIPTION_MAX_CHARS
    assert len(rec_text) <= pr.CALLOUT_MAX_CHARS
    # the raw strings are longer than the budgets above — truncation must
    # have actually kicked in, not just happened to fit.
    assert len(LONG_TITLE) > pr.TITLE_MAX_CHARS
    assert len(LONG_DESCRIPTION) > pr.DESCRIPTION_MAX_CHARS
    assert len(LONG_RECOMMENDATION) > pr.CALLOUT_MAX_CHARS


def test_table_rows_fit_within_the_space_reserved_before_the_callout_box():
    """5 rows at ROW_H starting at TABLE_TOP+HEADER_H must land exactly at
    (not past) the callout box's y — this is what keeps rows from ever
    visually colliding with the recommendation box below them."""
    rows_bottom = int(pr.TABLE_TOP) + int(pr.HEADER_H) + pr.ROWS_PER_PAGE * int(pr.ROW_H)
    callout_y = 4160520
    assert rows_bottom <= callout_y
